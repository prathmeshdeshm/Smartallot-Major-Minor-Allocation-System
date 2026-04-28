from django.test import TestCase
from django.utils import timezone
from django.contrib.auth.models import User
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from io import BytesIO
from openpyxl import Workbook

from .models import (
	Student, MinorBranch, OpenElective,
	MinorPreference, DoubleMinorPreference, OEPreference,
	MinorAllocation, DoubleMinorAllocation, OEAllocation,
	EligibilityRule, OEEligibilityRule, PreferenceWindow,
	StudentImport, ImportedStudent
)
from .utils import (
	run_minor1_allocation, run_minor2_allocation, run_oe_allocation,
	sync_existing_student_departments,
	import_students_from_excel,
	import_student_results_from_excel,
)


class AllocationLogicTests(TestCase):
	def setUp(self):
		# Ensure separate preference windows are open for saving preferences
		now = timezone.now()
		PreferenceWindow.objects.create(
			preference_type='minor',
			name="Minor Test Window",
			start_at=now - timezone.timedelta(days=1),
			end_at=now + timezone.timedelta(days=1),
			is_active=True,
		)
		PreferenceWindow.objects.create(
			preference_type='oe',
			name="OE Test Window",
			start_at=now - timezone.timedelta(days=1),
			end_at=now + timezone.timedelta(days=1),
			is_active=True,
		)

		# Create users + students with different percentages and departments
		self.u1 = User.objects.create_user(username='s1')
		self.u2 = User.objects.create_user(username='s2')
		self.u3 = User.objects.create_user(username='s3')

		self.s1 = Student.objects.create(
			user=self.u1,
			name='Alice',
			roll_no='R1',
			department='CSE',
			percentage=90,
			marks=90,
			email='a@example.com',
			is_validated=True,
			academic_status='CLEAR',
		)
		self.s2 = Student.objects.create(
			user=self.u2,
			name='Bob',
			roll_no='R2',
			department='IT',
			percentage=85,
			marks=85,
			email='b@example.com',
			is_validated=True,
			academic_status='CLEAR',
		)
		self.s3 = Student.objects.create(
			user=self.u3,
			name='Cara',
			roll_no='R3',
			department='ECE',
			percentage=70,
			marks=70,
			email='c@example.com',
			is_validated=True,
			academic_status='CLEAR',
		)

		# Create minor branches
		self.mb1 = MinorBranch.objects.create(name='AI', capacity=1, offering_dept='MECH')
		self.mb2 = MinorBranch.objects.create(name='IOT', capacity=2, offering_dept='EEE')

		# Create open electives
		self.oe1 = OpenElective.objects.create(name='Data Mining', capacity=1, offering_dept='MECH')
		self.oe2 = OpenElective.objects.create(name='AR-VR', capacity=2, offering_dept='EEE')

	def test_minor1_merit_and_preference(self):
		"""Higher percentage student gets limited capacity branch by priority."""
		# Both want AI first, IOT second
		MinorPreference.objects.create(student=self.s1, minor_branch=self.mb1, priority=1)
		MinorPreference.objects.create(student=self.s1, minor_branch=self.mb2, priority=2)
		MinorPreference.objects.create(student=self.s2, minor_branch=self.mb1, priority=1)
		MinorPreference.objects.create(student=self.s2, minor_branch=self.mb2, priority=2)

		run_minor1_allocation()

		a1 = MinorAllocation.objects.filter(student=self.s1).first()
		a2 = MinorAllocation.objects.filter(student=self.s2).first()

		self.assertIsNotNone(a1)
		self.assertEqual(a1.minor_branch, self.mb1)  # s1 (90%) gets AI
		# s2 should get IOT as fallback
		self.assertIsNotNone(a2)
		self.assertIn(a2.minor_branch, [self.mb2])

	def test_capacity_not_exceeded(self):
		"""No branch exceeds capacity across Minor 1 allocations."""
		# All three want AI, capacity=1
		for s in [self.s1, self.s2, self.s3]:
			MinorPreference.objects.create(student=s, minor_branch=self.mb1, priority=1)
		run_minor1_allocation()
		count_ai = MinorAllocation.objects.filter(minor_branch=self.mb1).count()
		self.assertLessEqual(count_ai, self.mb1.capacity)

	def test_ineligible_student_not_allocated(self):
		"""FAILED_REASSESSMENT + has_backlog students do not get allocated."""
		self.s3.has_backlog = True
		self.s3.academic_status = 'FAILED_REASSESSMENT'
		self.s3.save()
		MinorPreference.objects.create(student=self.s3, minor_branch=self.mb2, priority=1)
		run_minor1_allocation()
		self.assertFalse(MinorAllocation.objects.filter(student=self.s3).exists())

	def test_minor2_not_same_as_minor1(self):
		"""Minor 2 allocation must not equal Minor 1 branch."""
		# s1: Minor1 -> AI (capacity 1)
		MinorPreference.objects.create(student=self.s1, minor_branch=self.mb1, priority=1)
		run_minor1_allocation()

		# Minor 2 preferences include same branch and another
		DoubleMinorPreference.objects.create(student=self.s1, minor_branch=self.mb1, priority=1)
		DoubleMinorPreference.objects.create(student=self.s1, minor_branch=self.mb2, priority=2)
		run_minor2_allocation()

		m1 = MinorAllocation.objects.get(student=self.s1).minor_branch
		m2 = DoubleMinorAllocation.objects.get(student=self.s1).minor_branch
		self.assertNotEqual(m1, m2)

	def test_eligibility_rules_enforced(self):
		"""MIN_PERCENTAGE and DEPARTMENT_BLOCK rules are respected."""
		# mb2 requires >= 80, and blocks IT
		EligibilityRule.objects.create(branch=self.mb2, rule_type='MIN_PERCENTAGE', value={'min_percentage': 80}, is_active=True)
		EligibilityRule.objects.create(branch=self.mb2, rule_type='DEPARTMENT_BLOCK', value={'blocked_departments': ['IT']}, is_active=True)

		# s2(IT,85) blocked by department; s3(ECE,70) blocked by percentage
		MinorPreference.objects.create(student=self.s2, minor_branch=self.mb2, priority=1)
		MinorPreference.objects.create(student=self.s3, minor_branch=self.mb2, priority=1)
		run_minor1_allocation()

		self.assertFalse(MinorAllocation.objects.filter(student=self.s2, minor_branch=self.mb2).exists())
		self.assertFalse(MinorAllocation.objects.filter(student=self.s3, minor_branch=self.mb2).exists())

	def test_open_elective_allocation(self):
		"""OE allocation respects capacity and preferences."""
		OEPreference.objects.create(student=self.s1, oe_subject=self.oe1, priority=1)
		OEPreference.objects.create(student=self.s2, oe_subject=self.oe1, priority=1)
		OEPreference.objects.create(student=self.s2, oe_subject=self.oe2, priority=2)
		run_oe_allocation()

		oe1_count = OEAllocation.objects.filter(oe_subject=self.oe1).count()
		self.assertLessEqual(oe1_count, self.oe1.capacity)
		# Higher-merit s1 should win oe1; s2 should fall back to oe2
		a1 = OEAllocation.objects.filter(student=self.s1).first()
		a2 = OEAllocation.objects.filter(student=self.s2).first()
		self.assertEqual(a1.oe_subject, self.oe1)
		self.assertEqual(a2.oe_subject, self.oe2)


class PreferenceWindowSelectionTests(TestCase):
	def setUp(self):
		self.now = timezone.now()

		self.admin = User.objects.create_superuser(
			username='admin1',
			email='admin1@example.com',
			password='pass1234'
		)

		student_user = User.objects.create_user(
			username='student1',
			email='student1@example.com',
			password='pass1234'
		)
		self.student = Student.objects.create(
			user=student_user,
			name='Student One',
			roll_no='TST001',
			department='CSE',
			percentage=80,
			marks=80,
			email='student.one@example.com',
			is_validated=True,
		)

	def _create_minor_windows_with_conflicting_dates(self):
		active_recent = PreferenceWindow.objects.create(
			preference_type='minor',
			name='Recent Active Minor Window',
			start_at=self.now - timezone.timedelta(hours=1),
			end_at=self.now + timezone.timedelta(hours=1),
			is_active=True,
		)
		PreferenceWindow.objects.create(
			preference_type='minor',
			name='Future Inactive Minor Window',
			start_at=self.now + timezone.timedelta(days=7),
			end_at=self.now + timezone.timedelta(days=8),
			is_active=False,
		)
		return active_recent

	def test_admin_dashboard_prefers_active_window_over_later_inactive_window(self):
		active_recent = self._create_minor_windows_with_conflicting_dates()

		self.client.force_login(self.admin)
		response = self.client.get(reverse('admin_dashboard'))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['minor_window'].id, active_recent.id)

	def test_student_dashboard_uses_active_window_for_open_status(self):
		active_recent = self._create_minor_windows_with_conflicting_dates()

		session = self.client.session
		session['validated_student_id'] = self.student.id
		session['student_verification_complete'] = True
		session.save()

		response = self.client.get(reverse('student_dashboard'))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['minor_window'].id, active_recent.id)
		self.assertTrue(response.context['minor_window_is_open'])
		self.assertEqual(response.context['minor_window_status'], 'open')

	def test_admin_can_close_minor_window(self):
		active_window = PreferenceWindow.objects.create(
			preference_type='minor',
			name='Minor Window To Close',
			start_at=self.now - timezone.timedelta(hours=1),
			end_at=self.now + timezone.timedelta(hours=1),
			is_active=True,
		)

		self.client.force_login(self.admin)
		response = self.client.post(reverse('admin_dashboard'), {
			'close_window': '1',
			'preference_type': 'minor',
		})

		self.assertEqual(response.status_code, 302)
		active_window.refresh_from_db()
		self.assertFalse(active_window.is_active)

	def test_admin_can_close_oe_window(self):
		active_window = PreferenceWindow.objects.create(
			preference_type='oe',
			name='OE Window To Close',
			start_at=self.now - timezone.timedelta(hours=1),
			end_at=self.now + timezone.timedelta(hours=1),
			is_active=True,
		)

		self.client.force_login(self.admin)
		response = self.client.post(reverse('admin_dashboard'), {
			'close_window': '1',
			'preference_type': 'oe',
		})

		self.assertEqual(response.status_code, 302)
		active_window.refresh_from_db()
		self.assertFalse(active_window.is_active)


class StudentValidationReentryTests(TestCase):
	def setUp(self):
		self.user = User.objects.create_user(username='reentry-user')
		self.student = Student.objects.create(
			user=self.user,
			name='Reentry Student',
			roll_no='REN001',
			department='CSE',
			percentage=75,
			marks=450,
			email='reentry@student.com',
			is_validated=True,
		)

	def test_validate_view_allows_reentry_after_session_verified(self):
		session = self.client.session
		session['validated_student_id'] = self.student.id
		session['student_verification_complete'] = True
		session['validation_step_passed'] = False
		session.save()

		response = self.client.get(reverse('validate_student_form'))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['step'], 1)


class StudentDepartmentImportSyncTests(TestCase):
	def setUp(self):
		self.admin = User.objects.create_superuser(
			username='admin-sync',
			email='admin-sync@example.com',
			password='pass1234'
		)

		u1 = User.objects.create_user(username='sync-s1')
		u2 = User.objects.create_user(username='sync-s2')

		self.student1 = Student.objects.create(
			user=u1,
			name='Student Sync One',
			roll_no='SYNC001',
			department='IT',
			percentage=70,
			marks=70,
			email='sync1@example.com',
		)
		self.student2 = Student.objects.create(
			user=u2,
			name='Student Sync Two',
			roll_no='SYNC002',
			department='ECE',
			percentage=72,
			marks=72,
			email='sync2@example.com',
		)

		self.import_batch = StudentImport.objects.create(
			file=SimpleUploadedFile(
				'students.xlsx',
				b'fake-xlsx-content',
				content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
			),
			imported_by=self.admin,
			total_records=2,
			successful_records=2,
			failed_records=0,
		)

	def test_sync_updates_department_from_imported_major_branch(self):
		ImportedStudent.objects.create(
			import_batch=self.import_batch,
			full_name='Student Sync One',
			roll_no='SYNC001',
			marks=80,
			percentage=0,
			major_branch='CSE',
		)

		updated = sync_existing_student_departments(self.import_batch)
		self.student1.refresh_from_db()

		self.assertEqual(updated, 1)
		self.assertEqual(self.student1.department, 'CSE')

	def test_sync_ignores_general_major_branch(self):
		ImportedStudent.objects.create(
			import_batch=self.import_batch,
			full_name='Student Sync Two',
			roll_no='SYNC002',
			marks=81,
			percentage=0,
			major_branch='GENERAL',
		)

		updated = sync_existing_student_departments(self.import_batch)
		self.student2.refresh_from_db()

		self.assertEqual(updated, 0)
		self.assertEqual(self.student2.department, 'ECE')


class ImportedRecordsReportTests(TestCase):
	def setUp(self):
		self.admin = User.objects.create_superuser(
			username='admin-import-report',
			email='admin-import-report@example.com',
			password='pass1234'
		)

		excel_batch = StudentImport.objects.create(
			file=SimpleUploadedFile(
				'excel_students.xlsx',
				b'xlsx-content',
				content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
			),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)

		tr_batch = StudentImport.objects.create(
			file=SimpleUploadedFile(
				'tr_students.pdf',
				b'%PDF-1.4 fake',
				content_type='application/pdf'
			),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)

		ImportedStudent.objects.create(
			import_batch=excel_batch,
			full_name='Excel Student',
			roll_no='EX001',
			marks=75,
			percentage=0,
			major_branch='CSE',
		)

		ImportedStudent.objects.create(
			import_batch=tr_batch,
			full_name='TR Student',
			roll_no='TR001',
			marks=76,
			percentage=0,
			major_branch='IT',
		)

	def test_imported_records_report_shows_excel_and_tr_data(self):
		self.client.force_login(self.admin)
		response = self.client.get(reverse('imported_records_report'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Excel Student')
		self.assertContains(response, 'TR Student')

	def test_imported_records_report_source_filter_excel(self):
		self.client.force_login(self.admin)
		response = self.client.get(reverse('imported_records_report'), {'source': 'excel'})

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Excel Student')
		self.assertNotContains(response, 'TR Student')


class ExcelImportBranchPropagationTests(TestCase):
	def setUp(self):
		self.admin = User.objects.create_superuser(
			username='admin-excel-branch',
			email='admin-excel-branch@example.com',
			password='pass1234'
		)

		u1 = User.objects.create_user(username='excel-branch-1')
		u2 = User.objects.create_user(username='excel-branch-2')

		self.student1 = Student.objects.create(
			user=u1,
			name='Alpha Student',
			roll_no='2023001',
			department='IT',
			percentage=70,
			marks=420,
			email='alpha@student.com',
		)

		self.student2 = Student.objects.create(
			user=u2,
			name='Beta Student',
			roll_no='2023002',
			department='ECE',
			percentage=71,
			marks=430,
			email='beta@student.com',
		)

	def _build_excel_file(self):
		wb = Workbook()
		ws = wb.active
		ws.append(['Roll No', 'Name of Students', 'Grand Total', 'Branch'])
		ws.append([2023001, 'Alpha Student', 455, 'CSE'])
		ws.append([2023002.0, 'Beta Student', 448, None])

		buffer = BytesIO()
		wb.save(buffer)
		buffer.seek(0)

		return SimpleUploadedFile(
			'students.xlsx',
			buffer.read(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
		)

	def _build_excel_file_with_branch(self, branch_code):
		wb = Workbook()
		ws = wb.active
		ws.append(['Roll No', 'Name of Students', 'Grand Total', 'Branch'])
		ws.append([2023001, 'Alpha Student', 460, branch_code])
		ws.append([2023002.0, 'Beta Student', 452, None])

		buffer = BytesIO()
		wb.save(buffer)
		buffer.seek(0)

		return SimpleUploadedFile(
			f'students_{branch_code.lower()}.xlsx',
			buffer.read(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
		)

	def test_excel_cse_branch_sets_major_branch_and_student_department(self):
		excel_file = self._build_excel_file()

		import_batch, errors, synced_students = import_students_from_excel(excel_file, self.admin)

		self.assertIsNotNone(import_batch)
		self.assertEqual(errors, [])
		self.assertEqual(synced_students, 2)

		imported_rows = ImportedStudent.objects.filter(import_batch=import_batch).order_by('roll_no')
		self.assertEqual(imported_rows.count(), 2)
		self.assertTrue(all(row.major_branch == 'CSE' for row in imported_rows))

		self.student1.refresh_from_db()
		self.student2.refresh_from_db()
		self.assertEqual(self.student1.department, 'CSE')
		self.assertEqual(self.student2.department, 'CSE')

	def test_excel_cse_branch_updates_same_roll_imported_history(self):
		old_batch = StudentImport.objects.create(
			file=SimpleUploadedFile(
				'old.xlsx',
				b'old-content',
				content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
			),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)
		ImportedStudent.objects.create(
			import_batch=old_batch,
			full_name='Alpha Student',
			roll_no='2023001',
			marks=401,
			percentage=0,
			major_branch='IT',
		)

		excel_file = self._build_excel_file()
		import_students_from_excel(excel_file, self.admin)

		history_values = set(
			ImportedStudent.objects.filter(roll_no='2023001').values_list('major_branch', flat=True)
		)
		self.assertEqual(history_values, {'CSE'})

	def test_student_dashboard_shows_updated_department_after_import(self):
		excel_file = self._build_excel_file()
		import_students_from_excel(excel_file, self.admin)

		session = self.client.session
		session['validated_student_id'] = self.student1.id
		session['student_verification_complete'] = True
		session.save()

		response = self.client.get(reverse('student_dashboard'))
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['student'].department, 'CSE')

	def test_reimport_same_rolls_keeps_count_and_updates_department(self):
		first_file = self._build_excel_file_with_branch('CSE')
		import_students_from_excel(first_file, self.admin)

		self.assertEqual(ImportedStudent.objects.count(), 2)

		second_file = self._build_excel_file_with_branch('IT')
		import_students_from_excel(second_file, self.admin)

		self.assertEqual(ImportedStudent.objects.count(), 2)
		self.assertEqual(ImportedStudent.objects.filter(roll_no='2023001').count(), 1)
		self.assertEqual(ImportedStudent.objects.filter(roll_no='2023002').count(), 1)
		self.assertEqual(
			set(ImportedStudent.objects.values_list('major_branch', flat=True)),
			{'IT'}
		)

		self.student1.refresh_from_db()
		self.student2.refresh_from_db()
		self.assertEqual(self.student1.department, 'IT')
		self.assertEqual(self.student2.department, 'IT')


class LatestImportRefreshIntegrationTests(TestCase):
	def setUp(self):
		self.admin = User.objects.create_superuser(
			username='admin-latest-refresh',
			email='admin-latest-refresh@example.com',
			password='pass1234'
		)

		self.user = User.objects.create_user(username='refresh-student')
		self.student = Student.objects.create(
			user=self.user,
			name='Refresh Student',
			roll_no='2023555',
			department='IT',
			percentage=72,
			marks=430,
			email='refresh@student.com',
		)

	def _build_results_excel(self):
		wb = Workbook()
		ws = wb.active
		ws.append([
			'SR', 'Roll No', 'Student Name',
			'DBMS Th', 'DBMS Int', 'DBMS Total',
			'TOC Th', 'TOC Int', 'TOC Total',
			'SE Th', 'SE Int', 'SE Total',
			'DSS Th', 'DSS Int', 'DSS Total',
			'DM Th', 'DM Int', 'DM Total',
			'S&T Th', 'S&T Int', 'S&T Total',
			'Grand Total', 'Result', 'Branch'
		])
		ws.append([
			1, '2023555', 'Refresh Student',
			20, 20, 40,
			20, 20, 40,
			20, 20, 40,
			20, 20, 40,
			20, 20, 40,
			20, 20, 40,
			240, 'PASS', 'CSE'
		])

		buf = BytesIO()
		wb.save(buf)
		buf.seek(0)
		return SimpleUploadedFile(
			'results_with_branch.xlsx',
			buf.read(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
		)

	def test_results_import_updates_department_and_imported_branch(self):
		batch_old = StudentImport.objects.create(
			file=SimpleUploadedFile('old.xlsx', b'old', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)
		ImportedStudent.objects.create(
			import_batch=batch_old,
			full_name='Refresh Student',
			roll_no='2023555',
			marks=300,
			percentage=0,
			major_branch='IT',
		)

		successful, failed, errors = import_student_results_from_excel(self._build_results_excel(), self.admin)

		self.assertEqual(errors, [])
		self.assertEqual(successful, 1)
		self.assertEqual(failed, 0)

		self.student.refresh_from_db()
		self.assertEqual(self.student.department, 'CSE')
		self.assertEqual(
			ImportedStudent.objects.filter(roll_no='2023555').values_list('major_branch', flat=True).first(),
			'CSE'
		)

	def test_student_login_uses_latest_import_for_department(self):
		old_batch = StudentImport.objects.create(
			file=SimpleUploadedFile('old_batch.xlsx', b'old-batch', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)
		ImportedStudent.objects.create(
			import_batch=old_batch,
			full_name='Refresh Student',
			roll_no='2023555',
			marks=410,
			percentage=0,
			major_branch='IT',
		)

		new_batch = StudentImport.objects.create(
			file=SimpleUploadedFile('new_batch.xlsx', b'new-batch', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)
		ImportedStudent.objects.create(
			import_batch=new_batch,
			full_name='Refresh Student',
			roll_no='2023555',
			marks=420,
			percentage=0,
			major_branch='CSE',
		)

		response = self.client.post(reverse('student_login'), {
			'full_name': 'Refresh Student',
			'roll_no': '2023555',
		}, follow=True)

		self.assertEqual(response.status_code, 200)
		self.student.refresh_from_db()
		self.assertEqual(self.student.department, 'CSE')


class AdminDeleteModuleWorkflowTests(TestCase):
	def setUp(self):
		self.admin = User.objects.create_superuser(
			username='admin-delete-module',
			email='admin-delete-module@example.com',
			password='pass1234'
		)
		self.client.force_login(self.admin)

		student_user = User.objects.create_user(username='delete-student-user')
		self.student = Student.objects.create(
			user=student_user,
			name='Delete Student',
			roll_no='DEL001',
			department='CSE',
			percentage=75,
			marks=410,
			email='delete.student@example.com',
		)

		self.minor = MinorBranch.objects.create(name='AI', capacity=10, offering_dept='IT')
		self.oe = OpenElective.objects.create(name='Cloud Computing', capacity=10, offering_dept='ECE')

		self.import_batch = StudentImport.objects.create(
			file=SimpleUploadedFile(
				'delete_import.xlsx',
				b'xlsx-content',
				content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
			),
			imported_by=self.admin,
			total_records=1,
			successful_records=1,
			failed_records=0,
		)

		ImportedStudent.objects.create(
			import_batch=self.import_batch,
			full_name='Imported Delete Student',
			roll_no='IMP001',
			marks=390,
			percentage=0,
			major_branch='IT',
		)

	def test_courses_and_branches_delete_component_keeps_students(self):
		start_response = self.client.post(reverse('admin_delete_module'), {
			'action': 'start',
			'component': 'courses_branches',
		})
		self.assertEqual(start_response.status_code, 302)

		step1_response = self.client.post(reverse('admin_delete_confirm_step1'), {
			'confirm_step_1': 'on',
		})
		self.assertEqual(step1_response.status_code, 302)

		step2_response = self.client.post(reverse('admin_delete_confirm_step2'), {
			'confirm_step_2': 'on',
			'confirm_permanent': 'on',
			'confirmation_text': 'DELETE COURSES AND BRANCHES',
		}, follow=True)
		self.assertEqual(step2_response.status_code, 200)

		self.assertEqual(MinorBranch.objects.count(), 0)
		self.assertEqual(OpenElective.objects.count(), 0)
		self.assertEqual(Student.objects.count(), 1)
		self.assertEqual(ImportedStudent.objects.count(), 1)

	def test_students_delete_component_keeps_courses_and_branches(self):
		start_response = self.client.post(reverse('admin_delete_module'), {
			'action': 'start',
			'component': 'students',
		})
		self.assertEqual(start_response.status_code, 302)

		step1_response = self.client.post(reverse('admin_delete_confirm_step1'), {
			'confirm_step_1': 'on',
		})
		self.assertEqual(step1_response.status_code, 302)

		step2_response = self.client.post(reverse('admin_delete_confirm_step2'), {
			'confirm_step_2': 'on',
			'confirm_permanent': 'on',
			'confirmation_text': 'DELETE STUDENTS',
		}, follow=True)
		self.assertEqual(step2_response.status_code, 200)

		self.assertEqual(Student.objects.count(), 0)
		self.assertEqual(ImportedStudent.objects.count(), 0)
		self.assertEqual(StudentImport.objects.count(), 0)
		self.assertEqual(MinorBranch.objects.count(), 1)
		self.assertEqual(OpenElective.objects.count(), 1)

